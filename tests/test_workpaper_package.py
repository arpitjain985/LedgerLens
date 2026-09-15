import zipfile
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import Firm, Statement, Transaction, Invoice
from app.services.workpaper_package import build_workpaper_package


@pytest.fixture
def db_session():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


@pytest.fixture
def populated_firm(db_session):
    db_session.add(Firm(id="f1", name="Test Firm"))
    db_session.add(Statement(id="s1", firm_id="f1", filename="t.pdf"))
    db_session.add(Transaction(
        id="t1", statement_id="s1", description="Salary", amount=48000,
        txn_type="credit", predicted_ledger_head="SALARY",
    ))
    db_session.add(Transaction(
        id="t2", statement_id="s1", description="Big risky wire", amount=345000,
        txn_type="debit", predicted_ledger_head="SALES_RECEIPT",
        risk_level="High", risk_score=45, vendor="Random Corp",
    ))
    db_session.add(Invoice(
        id="i1", firm_id="f1", filename="inv.pdf", vendor_name="Raj Electricals",
        invoice_number="INV-001", taxable_amount=5000, cgst=450, sgst=450, igst=0, total_amount=5900,
    ))
    db_session.commit()
    return db_session


def test_package_contains_all_expected_files(populated_firm):
    zip_buffer = build_workpaper_package(populated_firm, "f1")
    z = zipfile.ZipFile(io.BytesIO(zip_buffer.getvalue()))
    assert set(z.namelist()) == {
        "Transaction_Register.xlsx", "Vendor_Summary.xlsx",
        "GST_Summary.xlsx", "Anomaly_Risk_Report.xlsx",
    }


def test_transaction_register_contains_all_transactions(populated_firm):
    zip_buffer = build_workpaper_package(populated_firm, "f1")
    z = zipfile.ZipFile(io.BytesIO(zip_buffer.getvalue()))
    wb = load_workbook(io.BytesIO(z.read("Transaction_Register.xlsx")))
    ws = wb["Workpaper"]
    assert ws.max_row - 1 == 2  # 2 transactions, minus 1 for header row


def test_gst_summary_totals_are_correct(populated_firm):
    zip_buffer = build_workpaper_package(populated_firm, "f1")
    z = zipfile.ZipFile(io.BytesIO(zip_buffer.getvalue()))
    wb = load_workbook(io.BytesIO(z.read("GST_Summary.xlsx")))
    ws = wb["GST Summary"]
    values = {row[0]: row[1] for row in ws.iter_rows(min_row=2, max_row=7, values_only=True)}
    assert values["Total Cgst"] == 450
    assert values["Total Sgst"] == 450
    assert values["Total Gst"] == 900


def test_anomaly_report_only_includes_flagged_transactions(populated_firm):
    """The normal Rs 48,000 salary credit must NOT appear -- only the
    flagged high-risk transaction."""
    zip_buffer = build_workpaper_package(populated_firm, "f1")
    z = zipfile.ZipFile(io.BytesIO(zip_buffer.getvalue()))
    wb = load_workbook(io.BytesIO(z.read("Anomaly_Risk_Report.xlsx")))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][2] == 345000  # amount column
    assert rows[0][6] == "High"  # risk level column


def test_empty_firm_produces_valid_but_empty_package(db_session):
    db_session.add(Firm(id="f2", name="Empty Firm"))
    db_session.commit()

    zip_buffer = build_workpaper_package(db_session, "f2")
    z = zipfile.ZipFile(io.BytesIO(zip_buffer.getvalue()))
    wb = load_workbook(io.BytesIO(z.read("Transaction_Register.xlsx")))
    ws = wb["Workpaper"]
    assert ws.max_row == 1  # header row only, no crash on empty data
