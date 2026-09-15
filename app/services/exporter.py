"""
Generates the actual deliverable: a CA-friendly Excel workpaper from
classified + reconciled transactions.
"""
from typing import List
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ANOMALY_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def build_workpaper(transactions: List[dict]) -> BytesIO:
    """
    transactions: list of dicts with keys matching schemas.TransactionOut
    Returns an in-memory .xlsx file ready to stream back to the client.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Workpaper"

    headers = [
        "Date", "Description", "Amount", "Type", "Ledger Head",
        "Confidence", "Method", "Reconciled", "Anomaly", "Anomaly Reason",
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for txn in transactions:
        row = [
            txn.get("txn_date"),
            txn.get("description"),
            txn.get("amount"),
            txn.get("txn_type"),
            txn.get("predicted_ledger_head"),
            txn.get("confidence"),
            txn.get("classification_method"),
            "Yes" if txn.get("reconciled") else "No",
            "Yes" if txn.get("is_anomaly") else "No",
            txn.get("anomaly_reason") or "",
        ]
        ws.append(row)
        if txn.get("is_anomaly"):
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col_idx).fill = ANOMALY_FILL

    # Auto-width columns roughly
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 4)

    # Summary sheet: total per ledger head — the part a CA actually needs for filing
    summary_ws = wb.create_sheet("Summary by Ledger Head")
    summary_ws.append(["Ledger Head", "Transaction Count", "Total Amount"])
    totals = {}
    counts = {}
    for txn in transactions:
        head = txn.get("predicted_ledger_head") or "UNCLASSIFIED"
        totals[head] = totals.get(head, 0) + (txn.get("amount") or 0)
        counts[head] = counts.get(head, 0) + 1
    for head in sorted(totals):
        summary_ws.append([head, counts[head], round(totals[head], 2)])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
