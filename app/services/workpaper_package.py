"""
Full CA workpaper package (Section 17): bundles multiple Excel reports into
a single ZIP, covering a firm's data across ALL statements/invoices
uploaded so far (not just one statement, unlike the single-statement export
in exporter.py which this module reuses for the Transaction Register sheet).

Honest scope: the spec lists 9 files including a PDF audit summary. This
implementation covers the 4 highest-value ones as real, working Excel
files: Transaction Register, Vendor Summary, GST Summary, and Anomaly/Risk
Report. Bank_Reconciliation.xlsx, Invoice_Reconciliation.xlsx, and the PDF
audit summary are NOT yet built — see PROGRESS.md. Four working, well-
formatted files are worth more than nine placeholder ones.
"""
from io import BytesIO
from typing import List
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import Transaction, Statement, Invoice
from app.services.exporter import build_workpaper
from app.services.vendor_intelligence import get_vendor_summary
from app.services.gst_summary import get_gst_summary

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RISK_FILL = {
    "Critical": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
    "High": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}


def build_workpaper_package(db: Session, firm_id: str) -> BytesIO:
    transactions = (
        db.query(Transaction)
        .join(Statement, Transaction.statement_id == Statement.id)
        .filter(Statement.firm_id == firm_id)
        .all()
    )
    txn_dicts = [_transaction_to_dict(t) for t in transactions]

    zip_buffer = BytesIO()
    with ZipFile(zip_buffer, "w", ZIP_DEFLATED) as zf:
        zf.writestr("Transaction_Register.xlsx", build_workpaper(txn_dicts).getvalue())
        zf.writestr("Vendor_Summary.xlsx", _build_vendor_summary_sheet(db, firm_id).getvalue())
        zf.writestr("GST_Summary.xlsx", _build_gst_summary_sheet(db, firm_id).getvalue())
        zf.writestr("Anomaly_Risk_Report.xlsx", _build_risk_report_sheet(txn_dicts).getvalue())

    zip_buffer.seek(0)
    return zip_buffer


def _transaction_to_dict(t: Transaction) -> dict:
    return {
        "txn_date": t.txn_date, "description": t.description, "amount": t.amount,
        "txn_type": t.txn_type, "predicted_ledger_head": t.corrected_ledger_head or t.predicted_ledger_head,
        "confidence": t.confidence, "classification_method": t.classification_method,
        "reconciled": t.reconciled, "is_anomaly": t.is_anomaly, "anomaly_reason": t.anomaly_reason,
        "risk_score": t.risk_score, "risk_level": t.risk_level, "vendor": t.vendor,
    }


def _new_workbook_with_header(title: str, headers: List[str]) -> tuple:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(headers[col_idx - 1]) + 4)
    return wb, ws


def _build_vendor_summary_sheet(db: Session, firm_id: str) -> BytesIO:
    vendors = get_vendor_summary(db, firm_id)
    headers = ["Vendor", "Total Payments", "Transaction Count", "Average Payment",
               "Highest Payment", "Lowest Payment", "First Payment", "Last Payment", "New Vendor", "Risk Score"]
    wb, ws = _new_workbook_with_header("Vendor Summary", headers)
    for v in vendors:
        ws.append([
            v["vendor_name"], v["total_payments"], v["transaction_count"], v["average_payment"],
            v["highest_payment"], v["lowest_payment"], v["first_payment_date"], v["last_payment_date"],
            "Yes" if v["is_new_vendor"] else "No", v["risk_score"],
        ])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _build_gst_summary_sheet(db: Session, firm_id: str) -> BytesIO:
    summary = get_gst_summary(db, firm_id)
    invoices = db.query(Invoice).filter(Invoice.firm_id == firm_id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "GST Summary"
    ws.append(["Metric", "Value"])
    for col_idx in (1, 2):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for key in ("invoice_count", "total_taxable_value", "total_cgst", "total_sgst", "total_igst", "total_gst", "total_invoice_value"):
        ws.append([key.replace("_", " ").title(), summary[key]])
    ws.append([])
    ws.append([summary["disclaimer"]])

    detail_ws = wb.create_sheet("Invoice Detail")
    headers = ["Invoice #", "Vendor", "Date", "Taxable Value", "CGST", "SGST", "IGST", "Total", "GSTIN"]
    detail_ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = detail_ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for inv in invoices:
        detail_ws.append([
            inv.invoice_number, inv.vendor_name, inv.invoice_date, inv.taxable_amount,
            inv.cgst, inv.sgst, inv.igst, inv.total_amount, inv.gstin_vendor,
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _build_risk_report_sheet(txn_dicts: List[dict]) -> BytesIO:
    flagged = [t for t in txn_dicts if t.get("is_anomaly") or (t.get("risk_level") in ("High", "Critical"))]
    headers = ["Date", "Description", "Amount", "Vendor", "Ledger Head", "Risk Score", "Risk Level", "Anomaly Reason"]
    wb, ws = _new_workbook_with_header("Anomaly & Risk Report", headers)
    for t in flagged:
        ws.append([
            t.get("txn_date"), t.get("description"), t.get("amount"), t.get("vendor"),
            t.get("predicted_ledger_head"), t.get("risk_score"), t.get("risk_level"), t.get("anomaly_reason") or "",
        ])
        fill = RISK_FILL.get(t.get("risk_level"))
        if fill:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col_idx).fill = fill
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
